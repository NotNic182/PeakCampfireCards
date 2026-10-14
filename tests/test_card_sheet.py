import importlib.util
import json
import shutil
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('card_sheet', ROOT / 'tools' / 'card_sheet.py')
sheet_tool = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sheet_tool)


class CardSheetTests(unittest.TestCase):
    """Export the card sheet, edit it like a player would, and apply it to temporary copies of the mod's files."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = Path(self.tmp.name)
        self.cards, self.plugin, self.xlsx = d / 'cards.json', d / 'Plugin.cs', d / 'card-sheet.xlsx'
        shutil.copy(sheet_tool.CARDS, self.cards)
        shutil.copy(sheet_tool.PLUGIN, self.plugin)
        self._saved = (sheet_tool.CARDS, sheet_tool.PLUGIN)
        sheet_tool.CARDS, sheet_tool.PLUGIN = self.cards, self.plugin
        sheet_tool.export(self.xlsx)

    def tearDown(self):
        sheet_tool.CARDS, sheet_tool.PLUGIN = self._saved
        self.tmp.cleanup()

    def edit(self, numbers=None, cards=None):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        for row in wb['Numbers'].iter_rows(min_row=2):
            key = (row[1].value, row[2].value)
            if numbers and key in numbers:
                row[5].value = numbers[key]
        for row in wb['Cards'].iter_rows(min_row=2):
            if cards and row[1].value in cards:
                for col, value in cards[row[1].value].items():
                    row[col].value = value
        wb.save(self.xlsx)

    def card(self, name):
        return next(c for c in json.loads(self.cards.read_text(encoding='utf-8'))['rows'] if c['name'] == name)

    def test_unchanged_sheet_changes_nothing(self):
        before = self.cards.read_text(encoding='utf-8')
        changes, warnings, numbers = sheet_tool.apply(self.xlsx)
        self.assertEqual([], changes)
        self.assertFalse(numbers)
        self.assertEqual(before, self.cards.read_text(encoding='utf-8'))

    def test_percent_seconds_and_points_convert_back(self):
        self.edit(numbers={('Daddy', 'Wall-grab reach'): 50, ('Bing BOING', 'Double jump cooldown'): 10,
                           ('Bing BOING', 'Drowsy added'): 5, ('Piton Fingers', 'Climbing stamina cost'): -20},
                  cards={'Daddy': {4: 'Grab walls from 50% farther away.'}})
        changes, warnings, numbers = sheet_tool.apply(self.xlsx)
        self.assertTrue(numbers)
        self.assertEqual(1.5, self.card('Daddy')['mods'][0][2])
        self.assertEqual(10, self.card('Bing BOING')['params']['cooldown'])
        self.assertAlmostEqual(0.05, self.card('Bing BOING')['params']['drowsy'])
        self.assertAlmostEqual(0.8, self.card('Piton Fingers')['mods'][0][2])
        self.assertEqual('Grab walls from 50% farther away.', self.card('Daddy')['positive'])
        self.assertFalse(any('Daddy' in w for w in warnings))
        self.assertTrue(any('Piton Fingers' in w for w in warnings))   # text still says 12%
        self.assertIn('CardBalanceVersion', self.plugin.read_text(encoding='utf-8'))

    def test_balance_version_goes_up_only_when_a_number_changes(self):
        import re
        version = lambda: int(re.search(r'CardBalanceVersion = (\d+);', self.plugin.read_text(encoding='utf-8')).group(1))
        start = version()
        self.edit(cards={'Daddy': {6: 'A flavor line.'}})
        sheet_tool.apply(self.xlsx)
        self.assertEqual(start, version())
        self.edit(numbers={('Daddy', 'Wall-grab reach'): 45})
        sheet_tool.apply(self.xlsx)
        self.assertEqual(start + 1, version())

    def test_text_name_rarity_and_empty_drawback(self):
        self.edit(cards={'Meditation': {1: 'Deep Breaths', 2: 'Rare', 5: ''}})
        sheet_tool.apply(self.xlsx)
        c = next(c for c in json.loads(self.cards.read_text(encoding='utf-8'))['rows'] if c['key'] == 'QuietBreathing')
        self.assertEqual('Deep Breaths', c['name'])
        self.assertEqual('rare', c['rarity'])
        self.assertEqual('', c['negative'])

    def test_check_only_writes_nothing(self):
        before = self.cards.read_text(encoding='utf-8')
        self.edit(numbers={('Daddy', 'Wall-grab reach'): 60})
        changes, _, _ = sheet_tool.apply(self.xlsx, check=True)
        self.assertEqual(1, len(changes))
        self.assertEqual(before, self.cards.read_text(encoding='utf-8'))

    def test_bad_values_are_refused_before_anything_is_written(self):
        before = self.cards.read_text(encoding='utf-8')
        for bad in (-150, 'lots', None):
            self.edit(numbers={('Daddy', 'Wall-grab reach'): bad})
            with self.assertRaises(ValueError):
                sheet_tool.apply(self.xlsx)
        self.assertEqual(before, self.cards.read_text(encoding='utf-8'))

    def test_bad_rarity_and_empty_name_are_refused(self):
        self.edit(cards={'Daddy': {2: 'Epic'}})
        with self.assertRaises(ValueError):
            sheet_tool.apply(self.xlsx)
        sheet_tool.export(self.xlsx)
        self.edit(cards={'Daddy': {1: ''}})
        with self.assertRaises(ValueError):
            sheet_tool.apply(self.xlsx)

    def test_difficulty_strengths_round_trip_and_validate(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        rows = {r[0].value: r for r in wb['Difficulty'].iter_rows(min_row=2)}
        self.assertEqual(100, rows['Tenderfoot – Peak'][1].value)   # all tiers default to 100% since the 0.7.1 rebalance
        rows['Ascent 1–3'][1].value = 80
        wb.save(self.xlsx)
        changes, _, numbers = sheet_tool.apply(self.xlsx)
        self.assertTrue(numbers)
        self.assertEqual(0.8, json.loads(self.cards.read_text(encoding='utf-8'))['difficultyStrength']['ascent1to3'])
        wb = load_workbook(self.xlsx)
        next(r for r in wb['Difficulty'].iter_rows(min_row=2) if r[0].value == 'Ascent 1–3')[1].value = 300
        wb.save(self.xlsx)
        with self.assertRaises(ValueError):
            sheet_tool.apply(self.xlsx)

    def test_every_card_and_number_is_in_the_sheet(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        data = json.loads(self.cards.read_text(encoding='utf-8'))
        self.assertEqual(len(data['rows']), wb['Cards'].max_row - 1)
        expected = sum(len(c['mods']) + len([p for p in c['params'] if p not in sheet_tool.HIDDEN_PARAMS]) for c in data['rows'])
        self.assertEqual(expected, wb['Numbers'].max_row - 1)
        self.assertNotIn('Key', [c.value for c in wb['Cards'][1]])   # no internal second name


if __name__ == '__main__':
    unittest.main()
