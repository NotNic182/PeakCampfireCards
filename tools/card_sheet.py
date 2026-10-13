#!/usr/bin/env python3
"""The card sheet: one simple workbook for editing every card, and the tool that puts it back into the mod.

  python tools/card_sheet.py export                  writes design/editable/card-sheet.xlsx from design/cards.json
  python tools/card_sheet.py apply <edited.xlsx>     writes the edits into design/cards.json (then regenerate and build)
  python tools/card_sheet.py apply <edited.xlsx> --check   only reports what would change

Cards tab: card number, name, rarity, benefit, drawback and flavor, exactly as shown in game.
Numbers tab: one row per number, in plain units (percent change, seconds, stamina points, metres...).
The card number and the hidden Ref column match each row back to the card; everything else is plain English.
Needs openpyxl.
"""
import argparse
import json
import math
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CARDS = ROOT / 'design' / 'cards.json'
STATS = ROOT / 'design' / 'stats.json'
CONDITIONS = ROOT / 'design' / 'conditions.json'
PLUGIN = ROOT / 'src' / 'CampfireCards' / 'Plugin.cs'
OUTPUT = ROOT / 'design' / 'editable' / 'card-sheet.xlsx'
RARITIES = ['common', 'uncommon', 'rare', 'legendary', 'mythic']

# ---------------------------------------------------------------------------------------------- units
PERCENT_CHANGE = '% change'          # +40 = 40% more, -12 = 12% less (the card multiplies by 1 + value/100)
POINTS = 'points'                    # stamina or status points: 25 = a quarter of the bar
PERCENT = '%'
SECONDS = 'seconds'
METRES = 'metres'
COUNT = 'cards'
STRENGTH = 'strength'

# flat (add) channels: unit and the factor from the stored value
ADD_UNITS = {'maxStamina': ('stamina points', 100.0), 'regenDelay': ('seconds (negative = sooner)', 1.0)}

# card parameters: (unit, factor from the stored value, plain label); per-card labels override
PARAMS = {
    'absorb': (PERCENT, 100.0, 'Share of the prevented Injury you take'),
    'bonus': (POINTS, 100.0, 'Bonus amount'),
    'bonusCards': (COUNT, 1.0, 'Bonus cards'),
    'bonusFire': (POINTS, 100.0, 'Yellow stamina at each later campfire'),
    'bonusNow': (POINTS, 100.0, 'Yellow stamina when picked'),
    'boost': (SECONDS, 1.0, 'Boost length'),
    'buildup': (SECONDS, 1.0, 'Time before it kicks in'),
    'cap': (POINTS, 100.0, 'Most Injury you take per hit'),
    'chance': (PERCENT, 100.0, 'Chance of a lucky roll'),
    'climbPenalty': (PERCENT_CHANGE, None, 'Possible penalty: climbing stamina cost'),
    'cooldown': (SECONDS, 1.0, 'Cooldown'),
    'curse': (POINTS, 100.0, 'Curse when picked'),
    'curseFire': (POINTS, 100.0, 'Curse at each later campfire'),
    'curseNow': (POINTS, 100.0, 'Curse when picked'),
    'defer': (PERCENT, 100.0, "Share of each fall's Injury that is delayed"),
    'delay': (SECONDS, 1.0, 'Delay before the rest lands'),
    'drowsy': (POINTS, 100.0, 'Drowsy added'),
    'duration': (SECONDS, 1.0, 'Power length'),
    'heal': (POINTS, 100.0, 'Injury removed'),
    'hungerPenalty': (PERCENT_CHANGE, None, 'Possible penalty: Hunger build-up'),
    'injury': (POINTS, 100.0, 'Injury'),
    'injuryPenalty': (PERCENT_CHANGE, None, 'Possible penalty: all Injury'),
    'interval': (SECONDS, 1.0, 'Time between powers'),
    'keep': (PERCENT, 100.0, 'Share of the delayed Injury that still lands'),
    'lowGravity': (STRENGTH, 1.0, 'Low Gravity strength (the mushroom is 3)'),
    'poison': (POINTS, 100.0, 'Poison added'),
    'radius': (METRES, 1.0, 'Range'),
    'regenPerSecond': (POINTS, 100.0, 'Stamina recovered per second'),
    'remove': (POINTS, 100.0, 'Amount removed'),
    'restore': (POINTS, 100.0, 'Stamina recovered'),
    'sprintPenalty': (PERCENT_CHANGE, None, 'Possible penalty: sprint stamina cost'),
    'stamina': (POINTS, 100.0, 'Stamina'),
    'threshold': (POINTS, 100.0, 'Threshold'),
}
CARD_LABELS = {
    ('CleansingRitual', 'bonus'): 'Extra status removed',
    ('HerosArrival', 'bonus'): 'Yellow stamina for both scouts',
    ('LoadedDice', 'bonus'): 'Yellow stamina on a lucky roll',
    ('BloodBattery', 'bonus'): 'Yellow stamina per use',
    ('PulseClimber', 'buildup'): 'Climbing time before the boost',
    ('MomentumFiend', 'buildup'): 'Sprint time before the speed-up',
    ('SecondSpring', 'cooldown'): 'Double jump cooldown',
    ('SecondWind', 'cooldown'): 'Cooldown',
    ('MomentumFiend', 'cooldown'): 'Time between Drowsy penalties',
    ('LoadedDice', 'injury'): 'Injury on an unlucky roll',
    ('BloodBattery', 'injury'): 'Injury per use',
    ('SafetyNet', 'injury'): 'Injury after the revive',
    ('EternalCampfire', 'injury'): 'Injury after the revive',
    ('LongArm', 'stamina'): 'Stamina cost to start a pull',
    ('EternalCampfire', 'stamina'): 'Stamina after the revive (at least)',
    ('SecondWind', 'threshold'): 'Works below this much stamina',
    ('EmergencyPurge', 'threshold'): 'Needs at least this much of the status',
    ('HungryHands', 'threshold'): 'Needs at least this much Injury',
    ('SecondWind', 'restore'): 'Stamina recovered',
    ('EmergencyPurge', 'remove'): 'Status removed',
    ('DoubleDown', 'bonusCards'): 'Bonus cards (Common to Rare)',
}
HIDDEN_PARAMS = {'item'}   # which item a card spawns: not a number to tweak
DIFFICULTY_TIERS = [('tenderfootPeak', 'Tenderfoot – Peak'), ('ascent1to3', 'Ascent 1–3'),
                    ('ascent4to6', 'Ascent 4–6'), ('ascent7to8', 'Ascent 7–8')]


def load(path):
    return json.loads(path.read_text(encoding='utf-8'))


def save_cards(sheet):
    CARDS.write_text(json.dumps(sheet, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


def tidy(v):
    """A stored number shown without float noise: 40.00000001 -> 40, 7.5 -> 7.5."""
    r = round(v, 4)
    return int(r) if abs(r - round(r)) < 1e-9 else r


def mod_unit(channel, op):
    if op == 'add':
        return ADD_UNITS.get(channel, ('', 1.0))
    return (PERCENT_CHANGE, None)


def to_display(stored, unit, factor):
    if unit == PERCENT_CHANGE:
        return tidy((stored - 1.0) * 100.0)
    return tidy(stored * factor)


def from_display(value, unit, factor):
    if unit == PERCENT_CHANGE:
        return round(1.0 + value / 100.0, 6)
    return round(value / factor, 6)


def who(target, card):
    radius = tidy(card['params'].get('radius', 6))
    return {'self': 'You', 'team': f'Teammates within {radius} m',
            'teamAndSelf': f'You and teammates within {radius} m'}.get(target, target)


def number_rows(card, stats, conditions):
    """Every editable number of one card: (ref, effect, when, who, value, unit)."""
    rows = []
    for i, (channel, op, value, cond, target) in enumerate(card['mods']):
        unit, factor = mod_unit(channel, op)
        label = stats[channel]['label'] if channel in stats else channel
        if op == 'set':
            label += ' (set, ignores other cards)'
        when = '' if cond == 'always' else 'while ' + conditions.get(cond, cond)
        rows.append((f'mod:{i}', label, when, who(target, card), to_display(value, unit, factor), unit))
    for name, value in card['params'].items():
        if name in HIDDEN_PARAMS:
            continue
        unit, factor, label = PARAMS.get(name, ('', 1.0, name))
        label = CARD_LABELS.get((card['key'], name), label)
        rows.append((f'param:{name}', label, '', '', to_display(value, unit, factor), unit))
    return rows


# ---------------------------------------------------------------------------------------------- export
def export(path=OUTPUT):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    sheet = load(CARDS)
    stats = {r['id']: r for r in load(STATS)['rows']}
    conditions = {r['id']: r['shown'] for r in load(CONDITIONS)['rows']}

    head_fill = PatternFill('solid', fgColor='2B1F17')
    head_font = Font(bold=True, color='FFFFFF')
    edit_fill = PatternFill('solid', fgColor='E3F1FB')
    ref_fill = PatternFill('solid', fgColor='F2F2F2')
    band_fill = PatternFill('solid', fgColor='EEF7FD')
    thin = Side(style='thin', color='D0D0D0')
    border = Border(bottom=thin)
    rarity_fills = {'common': 'D9D9D9', 'uncommon': 'C6EFCE', 'rare': 'C9DDF7', 'legendary': 'E4D2F7', 'mythic': 'FCE9A8'}

    wb = Workbook()
    guide = wb.active
    guide.title = 'How to edit'
    lines = [
        ('Campfire Cards card sheet', ''),
        ('Blue cells', 'Edit these. Grey cells are for reference; changes to them are ignored.'),
        ('Cards tab', 'Name, rarity, benefit, drawback and flavor exactly as the card shows them in game. Leave Drawback '
                      'empty for a card with no drawback. Keep {key} in key-press cards: the game fills in the key.'),
        ('Numbers tab', 'Every number a card uses, in plain units. "% change": +40 means 40% more, -12 means 12% less, '
                        '-100 means none at all. Points: 25 = a quarter of the stamina bar or a status bar.'),
        ('Text and numbers', 'The card text is what players read; the Numbers tab is what the game does. If you change a '
                             'number, change the text too (the apply step warns when the text still shows the old number).'),
        ('Difficulty tab', 'How strong card benefits are on each difficulty, in percent. The cards are written for Ascent 7–8 '
                           '(100%); on easier difficulties every benefit shrinks to this share, and the card text shows the '
                           'smaller number. Drawbacks and Mythics always stay at full strength.'),
        ('Card number', 'The # column matches each row to its card. Do not change it, and do not add or remove rows.'),
        ('Sending it back', 'Save as .xlsx and send it back; it is applied straight into the mod.'),
    ]
    for row in lines:
        guide.append(row)
    guide['A1'].font = Font(bold=True, size=14)
    guide.column_dimensions['A'].width = 20
    guide.column_dimensions['B'].width = 110
    for r in guide.iter_rows(min_row=2):
        r[0].font = Font(bold=True)
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical='top')

    # Cards tab
    cards = wb.create_sheet('Cards')
    headers = ['#', 'Card name', 'Rarity', 'Category', 'Benefit (shown on the card)', 'Drawback (blank = none)', 'Flavor (optional)']
    editable = {'Card name', 'Rarity', 'Benefit (shown on the card)', 'Drawback (blank = none)', 'Flavor (optional)'}
    cards.append(headers)
    for c in sheet['rows']:
        cards.append([c['id'], c['name'], c['rarity'].capitalize(), c['category'].capitalize(),
                      c['positive'], c['negative'], c.get('flavor', '')])
    widths = {'#': 5, 'Card name': 30, 'Rarity': 12, 'Category': 13, 'Benefit (shown on the card)': 60,
              'Drawback (blank = none)': 52, 'Flavor (optional)': 34}
    for i, h in enumerate(headers, 1):
        cards.column_dimensions[cards.cell(1, i).column_letter].width = widths[h]
    for row in cards.iter_rows(min_row=2):
        rarity = str(row[2].value).lower()
        for cell in row:
            h = headers[cell.column - 1]
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.fill = edit_fill if h in editable else ref_fill
            cell.border = border
            if isinstance(cell.value, str):
                cell.data_type = 's'
        row[2].fill = PatternFill('solid', fgColor=rarity_fills.get(rarity, 'FFFFFF'))
    rarity_list = DataValidation(type='list', formula1='"Common,Uncommon,Rare,Legendary,Mythic"', allow_blank=False)
    cards.add_data_validation(rarity_list)
    rarity_list.add(f'C2:C{cards.max_row}')

    # Numbers tab
    nums = wb.create_sheet('Numbers')
    nheaders = ['#', 'Card name', 'What it changes', 'When', 'Who', 'Value', 'Unit', 'Ref']
    nums.append(nheaders)
    band = False
    for c in sheet['rows']:
        band = not band
        for ref, effect, when, target, value, unit in number_rows(c, stats, conditions):
            nums.append([c['id'], c['name'], effect, when, target, value, unit, ref])
            for cell in nums[nums.max_row]:
                h = nheaders[cell.column - 1]
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.fill = edit_fill if h == 'Value' else (band_fill if band else ref_fill)
                cell.border = border
            nums.cell(nums.max_row, 6).number_format = '0.###'
    nwidths = {'#': 5, 'Card name': 30, 'What it changes': 42, 'When': 34, 'Who': 30, 'Value': 10, 'Unit': 26, 'Ref': 14}
    for i, h in enumerate(nheaders, 1):
        nums.column_dimensions[nums.cell(1, i).column_letter].width = nwidths[h]
    nums.column_dimensions['H'].hidden = True   # matches the row back to the card's number; not for editing

    diff = wb.create_sheet('Difficulty')
    dheaders = ['Difficulty', 'Card benefit strength (%)', 'What it means', 'Ref']
    diff.append(dheaders)
    tiers = sheet.get('difficultyStrength', {})
    for key, label in DIFFICULTY_TIERS:
        pct = tidy(tiers.get(key, 1.0) * 100.0)
        diff.append([label, pct, 'Benefits at this share of what the cards say; drawbacks and Mythics stay at 100%', key])
        for cell in diff[diff.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.fill = edit_fill if cell.column == 2 else ref_fill
            cell.border = border
    for col, w in zip('ABCD', (22, 24, 70, 14)):
        diff.column_dimensions[col].width = w
    diff.column_dimensions['D'].hidden = True
    diff.row_dimensions[1].height = 24
    for cell in diff[1]:
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical='center', wrap_text=True)

    for ws in (cards, nums):
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 24
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    wb.active = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------------------------- apply
def read_workbook(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    if 'Cards' not in wb.sheetnames or 'Numbers' not in wb.sheetnames:
        raise ValueError('The workbook needs its Cards and Numbers tabs')

    def rows(ws):
        values = list(ws.iter_rows(values_only=True))
        head = [str(h).strip() if h is not None else '' for h in values[0]]
        return [dict(zip(head, r)) for r in values[1:] if any(v not in (None, '') for v in r)]
    return rows(wb['Cards']), rows(wb['Numbers']), (rows(wb['Difficulty']) if 'Difficulty' in wb.sheetnames else [])


def text(v):
    return '' if v is None else str(v).strip()


def number(v, where):
    if isinstance(v, bool) or v is None or (isinstance(v, str) and not v.strip()):
        raise ValueError(f'{where}: Value is empty')
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        s = str(v).strip().replace('%', '').replace('+', '').replace(',', '.').strip()
        try:
            f = float(s)
        except ValueError:
            raise ValueError(f'{where}: "{v}" is not a number')
    if math.isnan(f) or math.isinf(f):
        raise ValueError(f'{where}: "{v}" is not a number')
    return f


def check_value(value, unit, where):
    if unit == PERCENT_CHANGE and value < -100:
        raise ValueError(f'{where}: {tidy(value)}% change would go below nothing (the lowest is -100)')
    if unit in (SECONDS, METRES, POINTS, STRENGTH, COUNT) and value < 0:
        raise ValueError(f'{where}: {tidy(value)} {unit} cannot be negative')
    if unit == PERCENT and not 0 <= value <= 100:
        raise ValueError(f'{where}: {tidy(value)}% must be between 0 and 100')
    if unit in (COUNT, STRENGTH) and abs(value - round(value)) > 1e-9:
        raise ValueError(f'{where}: {tidy(value)} must be a whole number')


def mentions(sentence, value):
    """Does a card text show this number (as 40, 40%, 40s, 7.5...)?"""
    shown = f'{tidy(abs(value))}'
    return re.search(r'(?<![\d.])' + re.escape(shown) + r'(?![\d])', sentence or '') is not None


def apply(path, check=False):
    sheet = load(CARDS)
    stats = {r['id']: r for r in load(STATS)['rows']}
    conditions = {r['id']: r['shown'] for r in load(CONDITIONS)['rows']}
    by_id = {c['id']: c for c in sheet['rows']}
    card_rows, num_rows, diff_rows = read_workbook(path)

    changes, warnings = [], []
    seen = set()
    for r in card_rows:
        where = f'Cards tab, card #{r.get("#")}'
        try:
            cid = int(r.get('#'))
        except (TypeError, ValueError):
            raise ValueError(f'Cards tab: a row has no card number (#): {r.get("Card name")!r}')
        if cid not in by_id:
            raise ValueError(f'{where}: there is no card #{cid}')
        if cid in seen:
            raise ValueError(f'{where}: card #{cid} appears twice')
        seen.add(cid)
        c = by_id[cid]
        name, benefit = text(r.get('Card name')), text(r.get('Benefit (shown on the card)'))
        if not name or not benefit:
            raise ValueError(f'{where}: the name and benefit cannot be empty')
        rarity = text(r.get('Rarity')).lower()
        if rarity not in RARITIES:
            raise ValueError(f'{where}: rarity "{r.get("Rarity")}" must be one of {", ".join(x.capitalize() for x in RARITIES)}')
        new = {'name': name, 'rarity': rarity, 'positive': benefit,
               'negative': text(r.get('Drawback (blank = none)')), 'flavor': text(r.get('Flavor (optional)'))}
        for field, value in new.items():
            if c.get(field, '') != value:
                changes.append(f'#{cid} {c["name"]}: {field} "{c.get(field, "")}" -> "{value}"')
                c[field] = value
    missing = set(by_id) - seen
    if missing:
        raise ValueError('Cards tab: missing card numbers ' + ', '.join(str(i) for i in sorted(missing)))

    numbers_changed = False
    expected = {(c['id'], ref): (unit, c) for c in sheet['rows'] for ref, _, _, _, _, unit in number_rows(c, stats, conditions)}
    seen_refs = set()
    for r in num_rows:
        try:
            cid = int(r.get('#'))
        except (TypeError, ValueError):
            raise ValueError(f'Numbers tab: a row has no card number (#): {r.get("What it changes")!r}')
        ref = text(r.get('Ref'))
        where = f'Numbers tab, card #{cid} "{r.get("What it changes")}"'
        if (cid, ref) not in expected:
            raise ValueError(f'{where}: this row does not match the card (keep the # and the hidden Ref column as they are)')
        if (cid, ref) in seen_refs:
            raise ValueError(f'{where}: appears twice')
        seen_refs.add((cid, ref))
        unit, c = expected[(cid, ref)]
        value = number(r.get('Value'), where)
        check_value(value, unit, where)
        kind, key = ref.split(':', 1)
        if kind == 'mod':
            i = int(key)
            mod = c['mods'][i]
            _, factor = mod_unit(mod[0], mod[1])
            old_display = to_display(mod[2], unit, factor)
            new_stored = from_display(value, unit, factor)
            if abs(new_stored - mod[2]) > 1e-6:
                mod[2] = tidy(new_stored) if isinstance(tidy(new_stored), int) else new_stored
                numbers_changed = True
                changes.append(f'#{cid} {c["name"]}: {r.get("What it changes")} {old_display} -> {tidy(value)} {unit}')
                if mentions(c['positive'] + ' ' + c['negative'], old_display) and not mentions(c['positive'] + ' ' + c['negative'], value):
                    warnings.append(f'#{cid} {c["name"]}: you changed "{r.get("What it changes")}" from {old_display} to {tidy(value)}, '
                                    f'but the card text still shows {tidy(abs(old_display))}')
        else:
            _, factor, _ = PARAMS.get(key, ('', 1.0, key))
            old = c['params'][key]
            old_display = to_display(old, unit, factor)
            new_stored = from_display(value, unit, factor)
            if abs(new_stored - old) > 1e-6:
                c['params'][key] = tidy(new_stored) if isinstance(tidy(new_stored), int) else new_stored
                numbers_changed = True
                changes.append(f'#{cid} {c["name"]}: {r.get("What it changes")} {old_display} -> {tidy(value)} {unit}')
                said = c['positive'] + ' ' + c['negative']
                shown_old = [old_display] + ([old_display / 60] if unit == SECONDS and old_display >= 60 else [])
                shown_new = [value] + ([value / 60] if unit == SECONDS and value >= 60 else [])
                if any(mentions(said, o) for o in shown_old) and not any(mentions(said, n) for n in shown_new):
                    warnings.append(f'#{cid} {c["name"]}: you changed "{r.get("What it changes")}" from {old_display} to {tidy(value)}, '
                                    f'but the card text still shows the old number')
    missing_refs = set(expected) - seen_refs
    if missing_refs:
        raise ValueError('Numbers tab: missing rows for cards ' + ', '.join(sorted({f'#{i}' for i, _ in missing_refs})))

    tiers = sheet.setdefault('difficultyStrength', {})
    labels = dict(DIFFICULTY_TIERS)
    for r in diff_rows:
        key = text(r.get('Ref'))
        if key not in labels:
            raise ValueError(f'Difficulty tab: row "{r.get("Difficulty")}" does not match a difficulty (keep the hidden Ref column)')
        where = f'Difficulty tab, {labels[key]}'
        value = number(r.get('Card benefit strength (%)'), where)
        if not 5 <= value <= 150:
            raise ValueError(f'{where}: {tidy(value)}% must be between 5 and 150')
        old = tiers.get(key, 1.0)
        if abs(value / 100.0 - old) > 1e-6:
            changes.append(f'Difficulty {labels[key]}: card benefit strength {tidy(old * 100)}% -> {tidy(value)}%')
            tiers[key] = tidy(round(value / 100.0, 4))
            numbers_changed = True

    if not check and changes:
        save_cards(sheet)
        if numbers_changed:
            bump_balance_version()
    return changes, warnings, numbers_changed


def bump_balance_version():
    """Saved configs keep old card numbers; a new balance version refreshes them once (Enabled choices are kept)."""
    src = PLUGIN.read_text(encoding='utf-8')
    m = re.search(r'public const int CardBalanceVersion = (\d+);', src)
    if not m:
        print('WARNING: CardBalanceVersion not found in Plugin.cs; bump it by hand')
        return
    n = int(m.group(1)) + 1
    src = src.replace(m.group(0), f'public const int CardBalanceVersion = {n};')
    src = re.sub(r'(/// 1 = 0\.3\.0[^\n]*?)\.\n', lambda x: x.group(1) + f', {n} = card sheet edit.\n', src, count=1)
    PLUGIN.write_text(src, encoding='utf-8')
    print(f'Balance version -> {n} (saved card numbers refresh once on next launch)')


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest='action', required=True)
    exp = sub.add_parser('export')
    exp.add_argument('--output', type=Path, default=OUTPUT)
    app = sub.add_parser('apply')
    app.add_argument('input', type=Path)
    app.add_argument('--check', action='store_true', help='report the changes without writing them')
    args = parser.parse_args()
    if args.action == 'export':
        print('Wrote', export(args.output))
        return
    try:
        changes, warnings, numbers_changed = apply(args.input, args.check)
    except ValueError as e:
        print('NOT APPLIED:', e)
        sys.exit(1)
    print(f'{len(changes)} change(s){" (check only, nothing written)" if args.check else ""}:')
    for c in changes:
        print('  ' + c)
    for w in warnings:
        print('WARNING: ' + w)
    if changes and not args.check:
        print('Next: python tools/gen.py, then tools/build.sh')


if __name__ == '__main__':
    main()
