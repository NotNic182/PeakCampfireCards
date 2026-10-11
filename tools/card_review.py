#!/usr/bin/env python3
"""Build the reviewed workbook and card-by-card Markdown table from the preserved review data.

Run after tools/card_catalog.py export --xlsx. Original requests remain literal text.
"""
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]


def main():
    report = json.loads((ROOT / 'docs/research/card-edit-review.json').read_text())
    current = {c['id']: c for c in json.loads((ROOT / 'design/cards.json').read_text())['rows']}
    # Read literal submitted text for every field, including fields the user left unchanged.
    # Falling back to shipped copy would misrepresent the request after a later wording pass.
    original = load_workbook(ROOT / 'design/reviews/card-editing-submitted.xlsx', read_only=True)
    values = iter(original['Cards'].values)
    headers = next(values)
    submitted = {int(row[0]): dict(zip(headers, row)) for row in values if row[0] is not None}
    original.close()
    wb = load_workbook(ROOT / 'design/editable/card-editing.xlsx')
    ws = wb.create_sheet('Review', 1)
    fields = ['id', 'key', 'status', 'requested_name', 'requested_benefit', 'requested_drawback', 'requested_flavor',
              'shipped_name', 'shipped_benefit', 'shipped_drawback', 'shipped_flavor', 'implementation_notes']
    ws.append(fields)
    table = ['| ID | Card | Status | Implementation notes |', '|---|---|---|---|']
    for row in report['cards']:
        card = current[row['id']]
        if card != row['shipped']:
            raise ValueError('Review data is stale for ' + card['key'])
        requested = {}
        for field, target in [('name', 'name'), ('benefit', 'positive'), ('drawback', 'negative'), ('flavor', 'flavor')]:
            requested[field] = submitted[row['id']].get(field) or ''
        status = 'Pending paired replacement' if row['status'] == 'pending' else 'Implemented — candidate'
        ws.append([row['id'], row['key'], status, requested['name'], requested['benefit'], requested['drawback'], requested['flavor'],
                   card['name'], card['positive'], card['negative'], card.get('flavor', ''), row['implementationNote']])
        esc = lambda value: str(value).replace('|', '\\|').replace('\n', '<br>')
        table.append('| ' + ' | '.join(map(esc, [row['id'], card['name'], status, row['implementationNote']])) + ' |')
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='173D45')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 8 if cell.value == 'id' else 65 if cell.value == 'implementation_notes' else 46
    for cells in ws.iter_rows(min_row=2):
        ws.row_dimensions[cells[0].row].height = 110
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.fill = PatternFill('solid', fgColor='FFF0CE' if cells[2].value.startswith('Pending') else 'F3F6F5')
            if isinstance(cell.value, str):
                cell.data_type = 's'
    wb.save(ROOT / 'design/editable/card-edit-review.xlsx')
    path = ROOT / 'docs/card-edit-review.md'
    text = path.read_text().split('<!-- reviewed-cards -->')[0]
    path.write_text(text + '<!-- reviewed-cards -->\n\n' + '\n'.join(table) + '\n')
    print('Reviewed %d cards; preserved requests and actual shipped definitions.' % len(report['cards']))


if __name__ == '__main__':
    main()
